from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.core.exceptions import ValidationFailed
from app.services.order_edit import (
    EditableField,
    apply_order_info_match,
    build_order_info_match_error_report,
    build_template_workbook,
    parse_requested_fields,
    preview_order_info_match,
)


class _ScalarsWrapper:
    def __init__(self, values):
        self._values = values

    def all(self):
        return self._values


class _RowsResult:
    def __init__(self, values):
        self._values = values

    def all(self):
        return self._values

    def scalars(self):
        return _ScalarsWrapper(self._values)


class _FakeSession:
    def __init__(self, responses, *, eu_countries=None):
        self.responses = list(responses)
        self.eu_countries = eu_countries
        self.committed = False
        self.statements = []

    async def get(self, model, identity):
        if self.eu_countries is None:
            return None
        return SimpleNamespace(eu_countries=self.eu_countries)

    async def execute(self, statement):
        self.statements.append(statement)
        if not self.responses:
            return _RowsResult([])
        return self.responses.pop(0)

    async def commit(self):
        self.committed = True


def _workbook_bytes(headers: list[str], rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def _base_parse_responses(extra=None):
    return [
        _RowsResult(["US"]),
        _RowsResult([]),
        _RowsResult(["ORDER-1"]),
        *(extra or []),
    ]


def test_parse_requested_fields_accepts_keys_and_labels() -> None:
    fields = parse_requested_fields("country_code,邮编")

    assert [field.key for field in fields] == ["country_code", "postal_code"]


@pytest.mark.parametrize(
    "field",
    [
        "shop_name",
        "order_platform",
        "order_total_amount",
        "fulfillment_channel",
        "purchase_date",
        "last_update_date",
        "marketplace_id",
        "Marketplace ID",
        "refund_status",
        "退款状态",
    ],
)
def test_parse_requested_fields_rejects_removed_edit_fields(field: str) -> None:
    with pytest.raises(ValidationFailed):
        parse_requested_fields(field)


def test_build_template_workbook_has_exact_header_order() -> None:
    fields = parse_requested_fields("country_code,postal_code")
    workbook = load_workbook(build_template_workbook(fields))

    assert [cell.value for cell in workbook.active[1]] == ["订单号", "国家", "邮编"]


@pytest.mark.asyncio
async def test_preview_order_info_match_reports_duplicate_and_missing_order() -> None:
    content = _workbook_bytes(
        ["订单号", "国家", "邮编"],
        [
            ["ORDER-1", "US", "90210"],
            ["ORDER-1", "US", "90210"],
            ["ORDER-404", "US", "90210"],
        ],
    )
    db = _FakeSession(_base_parse_responses())

    result = await preview_order_info_match(
        db,  # type: ignore[arg-type]
        fields=parse_requested_fields("country_code,postal_code"),
        content=content,
    )

    assert result.matched_order_count == 0
    assert result.matched_order_ids == []
    assert [error.row for error in result.errors] == [3, 4]
    assert "重复" in result.errors[0].message
    assert result.errors[1].message == "订单号不存在"


@pytest.mark.asyncio
async def test_order_info_match_error_report_preserves_content_and_appends_reasons() -> None:
    content = _workbook_bytes(
        ["订单号", "国家", "邮编"],
        [
            ["ORDER-1", "US", "90210"],
            ["ORDER-404", "XX", "10001"],
        ],
    )
    db = _FakeSession(_base_parse_responses())

    report = await build_order_info_match_error_report(
        db,  # type: ignore[arg-type]
        fields=parse_requested_fields("country_code,postal_code"),
        content=content,
    )

    sheet = load_workbook(report).active
    assert [sheet.cell(row=1, column=column).value for column in range(1, 5)] == [
        "订单号",
        "国家",
        "邮编",
        "错误原因",
    ]
    assert [sheet.cell(row=3, column=column).value for column in range(1, 4)] == [
        "ORDER-404",
        "XX",
        "10001",
    ]
    assert sheet.cell(row=2, column=4).value is None
    reason_cell = sheet.cell(row=3, column=4)
    assert reason_cell.value == "订单号：订单号不存在"
    assert reason_cell.font.color is not None
    assert reason_cell.font.color.rgb.endswith("FF0000")


@pytest.mark.asyncio
async def test_order_info_match_error_report_merges_multiple_errors_for_same_row() -> None:
    amount_field = EditableField("amount", "金额", "order_total_amount", "decimal")
    content = _workbook_bytes(
        ["订单号", "国家", "金额"],
        [["ORDER-1", "US - 美国", "not-a-number"]],
    )
    db = _FakeSession(_base_parse_responses())

    report = await build_order_info_match_error_report(
        db,  # type: ignore[arg-type]
        fields=[*parse_requested_fields("country_code"), amount_field],
        content=content,
    )

    sheet = load_workbook(report).active
    assert sheet.cell(row=2, column=4).value == (
        "国家：国家必须填写有效二字码；金额：订单金额必须是数值"
    )


@pytest.mark.asyncio
async def test_preview_order_info_match_requires_country_two_letter_code() -> None:
    content = _workbook_bytes(
        ["订单号", "国家"],
        [["ORDER-1", "US - 美国"]],
    )
    db = _FakeSession(_base_parse_responses())

    result = await preview_order_info_match(
        db,  # type: ignore[arg-type]
        fields=parse_requested_fields("country_code"),
        content=content,
    )

    assert result.matched_order_count == 0
    assert len(result.errors) == 1
    assert result.errors[0].field == "国家"
    assert result.errors[0].message == "国家必须填写有效二字码"


@pytest.mark.asyncio
async def test_preview_order_info_match_allows_blank_postal_code_to_clear() -> None:
    content = _workbook_bytes(
        ["订单号", "国家", "邮编"],
        [["ORDER-1", "US", None]],
    )
    db = _FakeSession(_base_parse_responses())

    result = await preview_order_info_match(
        db,  # type: ignore[arg-type]
        fields=parse_requested_fields("country_code,postal_code"),
        content=content,
    )

    assert result.matched_order_count == 1
    assert result.matched_order_ids == ["ORDER-1"]
    assert result.errors == []
    assert result.update_fields == ["国家", "邮编"]


@pytest.mark.asyncio
async def test_preview_order_info_match_maps_eu_member_country_from_config() -> None:
    content = _workbook_bytes(
        ["订单号", "国家"],
        [["ORDER-1", "DE"]],
    )
    db = _FakeSession(_base_parse_responses(), eu_countries=["DE", "FR"])

    result = await preview_order_info_match(
        db,  # type: ignore[arg-type]
        fields=parse_requested_fields("country_code"),
        content=content,
    )

    assert result.matched_order_count == 1
    assert result.matched_order_ids == ["ORDER-1"]
    assert result.errors == []


@pytest.mark.asyncio
async def test_apply_order_info_match_updates_all_packages_and_sets_manual_lock() -> None:
    header_1 = SimpleNamespace(amazon_order_id="ORDER-1", manual_edit_fields=None)
    header_2 = SimpleNamespace(amazon_order_id="ORDER-1", manual_edit_fields=["postal_code"])
    content = _workbook_bytes(
        ["订单号", "国家", "邮编"],
        [["ORDER-1", "US", "10001"]],
    )
    db = _FakeSession(_base_parse_responses(extra=[_RowsResult([header_1, header_2])]))

    result = await apply_order_info_match(
        db,  # type: ignore[arg-type]
        fields=parse_requested_fields("country_code,postal_code"),
        content=content,
        user_id=7,
    )

    assert result.updated_order_count == 2
    assert result.matched_order_ids == ["ORDER-1"]
    assert db.committed is True
    for header in (header_1, header_2):
        assert header.manual_edit_locked is True
        assert header.manual_edited_by == 7
        assert header.country_code == "US"
        assert header.original_country_code is None
        assert header.postal_code == "10001"
    assert header_2.manual_edit_fields == ["country_code", "postal_code"]


@pytest.mark.asyncio
async def test_apply_order_info_match_maps_eu_member_and_keeps_original_country() -> None:
    header = SimpleNamespace(amazon_order_id="ORDER-1", manual_edit_fields=None)
    content = _workbook_bytes(
        ["订单号", "国家"],
        [["ORDER-1", "DE"]],
    )
    db = _FakeSession(
        _base_parse_responses(extra=[_RowsResult([header])]),
        eu_countries=["DE", "FR"],
    )

    result = await apply_order_info_match(
        db,  # type: ignore[arg-type]
        fields=parse_requested_fields("country_code"),
        content=content,
        user_id=7,
    )

    assert result.updated_order_count == 1
    assert header.country_code == "EU"
    assert header.original_country_code == "DE"
    assert header.manual_edit_fields == ["country_code"]


@pytest.mark.asyncio
async def test_apply_order_info_match_clears_original_country_for_non_eu_country() -> None:
    header = SimpleNamespace(
        amazon_order_id="ORDER-1",
        manual_edit_fields=["country_code"],
        original_country_code="DE",
    )
    content = _workbook_bytes(
        ["订单号", "国家"],
        [["ORDER-1", "US"]],
    )
    db = _FakeSession(
        _base_parse_responses(extra=[_RowsResult([header])]),
        eu_countries=["DE", "FR"],
    )

    await apply_order_info_match(
        db,  # type: ignore[arg-type]
        fields=parse_requested_fields("country_code"),
        content=content,
        user_id=7,
    )

    assert header.country_code == "US"
    assert header.original_country_code is None
    assert header.manual_edit_fields == ["country_code"]


@pytest.mark.asyncio
async def test_apply_order_info_match_keeps_literal_eu_without_original_country() -> None:
    header = SimpleNamespace(
        amazon_order_id="ORDER-1",
        manual_edit_fields=None,
        original_country_code="DE",
    )
    content = _workbook_bytes(
        ["订单号", "国家"],
        [["ORDER-1", "EU"]],
    )
    db = _FakeSession(
        _base_parse_responses(extra=[_RowsResult([header])]),
        eu_countries=["DE", "FR"],
    )

    await apply_order_info_match(
        db,  # type: ignore[arg-type]
        fields=parse_requested_fields("country_code"),
        content=content,
        user_id=7,
    )

    assert header.country_code == "EU"
    assert header.original_country_code is None
    assert header.manual_edit_fields == ["country_code"]
