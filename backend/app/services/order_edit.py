"""Order manual edit and Excel information matching service."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from sqlalchemy import func, select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.countries import BUILTIN_COUNTRY_NAMES, country_label, normalize_observed_country_code
from app.core.exceptions import NotFound, ValidationFailed
from app.core.timezone import BEIJING, now_beijing
from app.models.country import CountryNameOverride
from app.models.order import ORDER_SOURCE_PACKAGE, OrderHeader
from app.models.shop import Shop
from app.schemas.data import (
    DataOrderPatch,
    OrderInfoMatchApplyOut,
    OrderInfoMatchError,
    OrderInfoMatchPreviewOut,
)

ORDER_NUMBER_HEADER = "订单号"
DATETIME_FORMAT = "%Y-%m-%d %H:%M:%S"
NEW_COUNTRY_RE = re.compile(r"^([A-Za-z]{2})\s*-\s*(.+)$")


@dataclass(frozen=True, slots=True)
class EditableField:
    key: str
    label: str
    model_field: str
    kind: str


EDITABLE_FIELDS: tuple[EditableField, ...] = (
    EditableField("shop_name", "店铺名称", "shop_name", "shop"),
    EditableField("order_platform", "平台", "order_platform", "platform"),
    EditableField("country_code", "国家", "country_code", "country"),
    EditableField("postal_code", "邮编", "postal_code", "text"),
    EditableField("marketplace_id", "Marketplace ID", "marketplace_id", "text"),
    EditableField("order_total_amount", "订单金额", "order_total_amount", "decimal"),
    EditableField("order_total_currency", "币种", "order_total_currency", "text"),
    EditableField("fulfillment_channel", "履约渠道", "fulfillment_channel", "text"),
    EditableField("purchase_date", "下单时间", "purchase_date", "datetime"),
    EditableField("last_update_date", "最后更新时间", "last_update_date", "datetime"),
    EditableField("refund_status", "退款状态", "refund_status", "text"),
)

FIELD_BY_KEY = {field.key: field for field in EDITABLE_FIELDS}
FIELD_BY_LABEL = {field.label: field for field in EDITABLE_FIELDS}
FIELD_BY_MODEL = {field.model_field: field for field in EDITABLE_FIELDS}


@dataclass(slots=True)
class ParsedWorkbook:
    updates_by_order: dict[str, dict[str, Any]]
    country_overrides: dict[str, str]
    matched_order_count: int
    update_fields: list[str]
    errors: list[OrderInfoMatchError]


def parse_requested_fields(raw_fields: str | list[str] | tuple[str, ...] | None) -> list[EditableField]:
    if raw_fields is None:
        parts: list[str] = []
    elif isinstance(raw_fields, str):
        parts = [part.strip() for part in raw_fields.split(",")]
    else:
        parts = []
        for item in raw_fields:
            parts.extend(str(item).split(","))
        parts = [part.strip() for part in parts]

    fields: list[EditableField] = []
    seen: set[str] = set()
    for part in parts:
        if not part:
            continue
        normalized = part.strip()
        field = FIELD_BY_KEY.get(normalized) or FIELD_BY_LABEL.get(normalized)
        if field is None:
            camel_to_snake = re.sub(r"(?<!^)(?=[A-Z])", "_", normalized).lower()
            field = FIELD_BY_KEY.get(camel_to_snake)
        if field is None:
            raise ValidationFailed(f"未知匹配字段：{part}")
        if field.key in seen:
            raise ValidationFailed(f"重复匹配字段：{field.label}")
        seen.add(field.key)
        fields.append(field)
    if not fields:
        raise ValidationFailed("请至少选择一个匹配字段")
    return fields


def build_template_workbook(fields: list[EditableField]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "订单信息匹配"
    sheet.append([ORDER_NUMBER_HEADER, *[field.label for field in fields]])
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


async def patch_order_header(
    db: AsyncSession,
    *,
    shop_id: str,
    amazon_order_id: str,
    package_sn: str | None,
    patch: DataOrderPatch,
    user_id: int,
) -> OrderHeader:
    updates = patch.model_dump(exclude_unset=True, by_alias=False)
    invalid = [key for key in updates if key not in FIELD_BY_MODEL]
    if invalid:
        raise ValidationFailed(f"不可编辑字段：{', '.join(invalid)}")
    if not updates:
        raise ValidationFailed("没有提交任何可更新字段")

    shop_map = await _load_shop_map(db)
    platform_options = await _load_platform_options(db)
    country_options = await _load_country_options(db)
    normalized, country_overrides, errors = _normalize_update_values(
        updates,
        fields=[FIELD_BY_MODEL[key] for key in updates],
        row_number=0,
        shop_map=shop_map,
        platform_options=platform_options,
        country_options=country_options,
        require_non_empty=False,
    )
    if errors:
        raise ValidationFailed("订单编辑校验失败", detail={"errors": [e.model_dump() for e in errors]})

    stmt = select(OrderHeader).where(
        OrderHeader.shop_id == shop_id,
        OrderHeader.amazon_order_id == amazon_order_id,
        OrderHeader.source == ORDER_SOURCE_PACKAGE,
    )
    if package_sn is not None:
        stmt = stmt.where(OrderHeader.package_sn == package_sn)
    else:
        stmt = stmt.order_by(OrderHeader.purchase_date.desc(), OrderHeader.id.desc()).limit(1)
    header = (await db.execute(stmt)).scalar_one_or_none()
    if header is None:
        raise NotFound(f"订单 {shop_id}/{amazon_order_id} 不存在")

    await _save_country_overrides(db, country_overrides)
    _apply_manual_updates(header, normalized, user_id=user_id)
    await db.commit()
    await db.refresh(header)
    return header


async def preview_order_info_match(
    db: AsyncSession,
    *,
    fields: list[EditableField],
    content: bytes,
) -> OrderInfoMatchPreviewOut:
    parsed = await _parse_match_workbook(db, fields=fields, content=content)
    return OrderInfoMatchPreviewOut(
        matched_order_count=parsed.matched_order_count,
        update_fields=parsed.update_fields,
        errors=parsed.errors,
    )


async def apply_order_info_match(
    db: AsyncSession,
    *,
    fields: list[EditableField],
    content: bytes,
    user_id: int,
) -> OrderInfoMatchApplyOut:
    parsed = await _parse_match_workbook(db, fields=fields, content=content)
    if parsed.errors:
        raise ValidationFailed(
            "订单信息匹配校验失败",
            detail={"errors": [error.model_dump() for error in parsed.errors]},
        )
    await _save_country_overrides(db, parsed.country_overrides)
    updated = 0
    for amazon_order_id, updates in parsed.updates_by_order.items():
        rows = (
            (
                await db.execute(
                    select(OrderHeader).where(
                        OrderHeader.amazon_order_id == amazon_order_id,
                        OrderHeader.source == ORDER_SOURCE_PACKAGE,
                    )
                )
            )
            .scalars()
            .all()
        )
        for row in rows:
            _apply_manual_updates(row, updates, user_id=user_id)
            updated += 1
    await db.commit()
    return OrderInfoMatchApplyOut(
        matched_order_count=parsed.matched_order_count,
        update_fields=parsed.update_fields,
        errors=[],
        updated_order_count=updated,
    )


async def _parse_match_workbook(
    db: AsyncSession,
    *,
    fields: list[EditableField],
    content: bytes,
) -> ParsedWorkbook:
    if not content:
        raise ValidationFailed("导入文件没有有效数据")
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:  # pragma: no cover - openpyxl exception types vary
        raise ValidationFailed("无法读取 Excel 文件") from exc

    sheet = workbook.active
    expected_headers = [ORDER_NUMBER_HEADER, *[field.label for field in fields]]
    header_values = [
        _string_cell(sheet.cell(row=1, column=column).value)
        for column in range(1, len(expected_headers) + 1)
    ]
    trailing_values = [
        _string_cell(sheet.cell(row=1, column=column).value)
        for column in range(len(expected_headers) + 1, sheet.max_column + 1)
    ]
    if header_values != expected_headers or any(trailing_values):
        raise ValidationFailed(
            "表头必须严格等于：订单号 + 勾选字段",
            detail={
                "expected": expected_headers,
                "actual": header_values + [value for value in trailing_values if value],
            },
        )

    shop_map = await _load_shop_map(db)
    platform_options = await _load_platform_options(db)
    country_options = await _load_country_options(db)
    all_order_ids = set(
        (
            await db.execute(
                select(OrderHeader.amazon_order_id)
                .where(OrderHeader.source == ORDER_SOURCE_PACKAGE)
                .distinct()
            )
        )
        .scalars()
        .all()
    )

    errors: list[OrderInfoMatchError] = []
    seen_orders: dict[str, int] = {}
    updates_by_order: dict[str, dict[str, Any]] = {}
    country_overrides: dict[str, str] = {}
    for row_number in range(2, sheet.max_row + 1):
        row_values = [
            sheet.cell(row=row_number, column=column).value
            for column in range(1, len(expected_headers) + 1)
        ]
        if all(_string_cell(value) == "" for value in row_values):
            continue
        amazon_order_id = _string_cell(row_values[0])
        if not amazon_order_id:
            errors.append(OrderInfoMatchError(row=row_number, field=ORDER_NUMBER_HEADER, message="订单号不能为空"))
            continue
        previous_row = seen_orders.get(amazon_order_id)
        if previous_row is not None:
            errors.append(
                OrderInfoMatchError(
                    row=row_number,
                    field=ORDER_NUMBER_HEADER,
                    message=f"订单号与第 {previous_row} 行重复",
                )
            )
            continue
        seen_orders[amazon_order_id] = row_number
        if amazon_order_id not in all_order_ids:
            errors.append(
                OrderInfoMatchError(row=row_number, field=ORDER_NUMBER_HEADER, message="订单号不存在")
            )
            continue

        raw_updates = {
            field.model_field: row_values[index + 1]
            for index, field in enumerate(fields)
        }
        normalized, row_country_overrides, row_errors = _normalize_update_values(
            raw_updates,
            fields=fields,
            row_number=row_number,
            shop_map=shop_map,
            platform_options=platform_options,
            country_options=country_options,
            require_non_empty=True,
        )
        errors.extend(row_errors)
        if not row_errors:
            updates_by_order[amazon_order_id] = normalized
            country_overrides.update(row_country_overrides)

    return ParsedWorkbook(
        updates_by_order=updates_by_order if not errors else {},
        country_overrides=country_overrides if not errors else {},
        matched_order_count=len(updates_by_order) if not errors else 0,
        update_fields=[field.label for field in fields],
        errors=errors,
    )


def _normalize_update_values(
    raw_updates: dict[str, Any],
    *,
    fields: list[EditableField],
    row_number: int,
    shop_map: dict[str, str],
    platform_options: set[str],
    country_options: dict[str, str],
    require_non_empty: bool,
) -> tuple[dict[str, Any], dict[str, str], list[OrderInfoMatchError]]:
    normalized: dict[str, Any] = {}
    country_overrides: dict[str, str] = {}
    errors: list[OrderInfoMatchError] = []
    for field in fields:
        raw_value = raw_updates.get(field.model_field)
        text_value = _string_cell(raw_value)
        if require_non_empty and not text_value:
            errors.append(OrderInfoMatchError(row=row_number, field=field.label, message="该字段不能为空"))
            continue
        if not require_non_empty and raw_value is None:
            normalized[field.model_field] = None
            continue
        try:
            if field.kind == "shop":
                canonical = shop_map.get(text_value)
                if canonical is None:
                    raise ValueError("店铺名称必须匹配现有店铺名称或店铺 ID")
                normalized[field.model_field] = canonical
            elif field.kind == "platform":
                if text_value not in platform_options:
                    raise ValueError("平台必须存在于当前已落库平台选项")
                normalized[field.model_field] = text_value
            elif field.kind == "country":
                code, name = _normalize_country_value(text_value, country_options)
                normalized[field.model_field] = code
                if name is not None:
                    country_overrides[code] = name
                    country_options[code] = country_label(code, {code: name})
            elif field.kind == "decimal":
                normalized[field.model_field] = _parse_decimal(raw_value)
            elif field.kind == "datetime":
                normalized[field.model_field] = _parse_datetime(raw_value)
            else:
                normalized[field.model_field] = text_value or None
        except ValueError as exc:
            errors.append(OrderInfoMatchError(row=row_number, field=field.label, message=str(exc)))
    return normalized, country_overrides, errors


def _normalize_country_value(value: str, options: dict[str, str]) -> tuple[str, str | None]:
    text = value.strip()
    code = normalize_observed_country_code(text)
    if code is not None and code in options:
        return code, None
    for option_code, label in options.items():
        if text == label:
            return option_code, None
    match = NEW_COUNTRY_RE.match(text)
    if not match:
        raise ValueError("国家必须使用已有代码/标签，或填写 XX - 中文名")
    new_code = normalize_observed_country_code(match.group(1))
    name = match.group(2).strip()
    if new_code is None or not name:
        raise ValueError("新国家必须填写有效的 XX - 中文名")
    if new_code in options:
        return new_code, None
    return new_code, name


def _parse_decimal(value: Any) -> Decimal:
    if isinstance(value, Decimal):
        return value
    text = _string_cell(value)
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("订单金额必须是数值") from exc


def _parse_datetime(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value.astimezone(BEIJING) if value.tzinfo else value.replace(tzinfo=BEIJING)
    text = _string_cell(value)
    try:
        return datetime.strptime(text, DATETIME_FORMAT).replace(tzinfo=BEIJING)
    except ValueError as exc:
        raise ValueError("日期必须为 YYYY-MM-DD HH:mm:ss") from exc


def _string_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime(DATETIME_FORMAT)
    return str(value).strip()


async def _load_shop_map(db: AsyncSession) -> dict[str, str]:
    rows = (await db.execute(select(Shop.id, Shop.name))).all()
    result: dict[str, str] = {}
    for shop_id, shop_name in rows:
        if shop_id:
            result[str(shop_id).strip()] = shop_name
        if shop_name:
            result[str(shop_name).strip()] = shop_name
    return result


async def _load_platform_options(db: AsyncSession) -> set[str]:
    platform_expr = func.trim(OrderHeader.order_platform)
    rows = (
        await db.execute(
            select(platform_expr)
            .where(OrderHeader.order_platform.is_not(None), platform_expr != "")
            .distinct()
        )
    ).scalars().all()
    return {platform for platform in rows if platform}


async def _load_country_options(db: AsyncSession) -> dict[str, str]:
    observed_rows = (
        await db.execute(
            select(OrderHeader.country_code)
            .where(OrderHeader.country_code.is_not(None), OrderHeader.country_code != "")
            .distinct()
        )
    ).scalars().all()
    override_rows = (
        (await db.execute(select(CountryNameOverride))).scalars().all()
    )
    overrides = {row.code.upper(): row.name for row in override_rows}
    codes = set(BUILTIN_COUNTRY_NAMES) | set(overrides)
    for raw in observed_rows:
        code = normalize_observed_country_code(raw)
        if code is not None:
            codes.add(code)
    return {code: country_label(code, overrides) for code in codes}


async def _save_country_overrides(db: AsyncSession, overrides: dict[str, str]) -> None:
    if not overrides:
        return
    now = now_beijing()
    stmt = pg_insert(CountryNameOverride).values(
        [
            {
                "code": code,
                "name": name,
                "created_at": now,
                "updated_at": now,
            }
            for code, name in sorted(overrides.items())
        ]
    )
    stmt = stmt.on_conflict_do_update(
        index_elements=["code"],
        set_={"name": stmt.excluded.name, "updated_at": now},
    )
    await db.execute(stmt)


def _apply_manual_updates(header: OrderHeader, updates: dict[str, Any], *, user_id: int) -> None:
    for field, value in updates.items():
        setattr(header, field, value)
    existing_fields = set(header.manual_edit_fields or [])
    existing_fields.update(updates)
    header.manual_edit_locked = True
    header.manual_edited_at = now_beijing()
    header.manual_edited_by = user_id
    header.manual_edit_fields = sorted(existing_fields)
