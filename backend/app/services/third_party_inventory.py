"""Third-party inventory import and aggregation helpers."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy import delete, func, select, update
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import ConflictError, NotFound, ValidationFailed
from app.models.third_party_inventory import (
    ThirdPartyInventoryCurrent,
    ThirdPartyInventoryImportBatch,
    ThirdPartyInventoryImportItem,
    ThirdPartyWarehouse,
)

REQUIRED_HEADERS = ("仓库", "SKU", "可用数", "待出库")
PREVIEW_ISSUE_LIMIT = 20
BEIJING = ZoneInfo("Asia/Shanghai")


@dataclass(frozen=True, slots=True)
class ParsedThirdPartyInventoryRow:
    source_row_no: int
    warehouse_name: str
    commodity_sku: str | None
    available: int | None
    reserved: int | None
    error_message: str | None


def _cell_text(value: Any) -> str:
    return str(value if value is not None else "").strip()


def _parse_quantity(value: Any, field_name: str) -> int:
    if value is None or str(value).strip() == "":
        return 0
    try:
        parsed = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field_name} 必须为非负整数") from exc
    if parsed < 0:
        raise ValueError(f"{field_name} 必须为非负整数")
    return parsed


def parse_third_party_inventory_workbook(content: bytes) -> list[ParsedThirdPartyInventoryRow]:
    """Parse the required Excel columns into validated row objects."""
    if not content:
        raise ValidationFailed("导入文件没有有效数据")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # pragma: no cover - openpyxl exception types vary
        raise ValidationFailed("导入文件无法解析，请上传 Excel 文件") from exc
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        raise ValidationFailed("导入文件没有有效数据") from exc

    headers = [_cell_text(cell) for cell in header_row]
    header_index = {header: idx for idx, header in enumerate(headers) if header}
    missing = [header for header in REQUIRED_HEADERS if header not in header_index]
    if missing:
        raise ValidationFailed(f"导入文件缺少必要列：{', '.join(missing)}")

    parsed_rows: list[ParsedThirdPartyInventoryRow] = []
    for row_no, row in enumerate(rows, start=2):
        if not any(_cell_text(value) for value in row):
            continue
        warehouse_name = _cell_text(row[header_index["仓库"]] if header_index["仓库"] < len(row) else None)
        commodity_sku = _cell_text(row[header_index["SKU"]] if header_index["SKU"] < len(row) else None)
        raw_available = row[header_index["可用数"]] if header_index["可用数"] < len(row) else None
        raw_reserved = row[header_index["待出库"]] if header_index["待出库"] < len(row) else None

        errors: list[str] = []
        available: int | None = None
        reserved: int | None = None
        if not warehouse_name:
            errors.append("仓库不能为空")
        if not commodity_sku:
            errors.append("SKU 不能为空")
        try:
            available = _parse_quantity(raw_available, "可用数")
        except ValueError as exc:
            errors.append(str(exc))
        try:
            reserved = _parse_quantity(raw_reserved, "待出库")
        except ValueError as exc:
            errors.append(str(exc))

        parsed_rows.append(
            ParsedThirdPartyInventoryRow(
                source_row_no=row_no,
                warehouse_name=warehouse_name,
                commodity_sku=commodity_sku or None,
                available=available if not errors else None,
                reserved=reserved if not errors else None,
                error_message="；".join(errors) if errors else None,
            )
        )
    if not parsed_rows:
        raise ValidationFailed("导入文件没有有效数据")
    return parsed_rows


async def create_import_preview(
    db: AsyncSession,
    *,
    filename: str,
    content: bytes,
    created_by: str | None,
) -> ThirdPartyInventoryImportBatch:
    parsed_rows = parse_third_party_inventory_workbook(content)
    valid_rows = [row for row in parsed_rows if row.error_message is None]
    valid_warehouse_names = sorted({row.warehouse_name for row in valid_rows})
    existing_rows = (
        await db.execute(
            select(ThirdPartyWarehouse).where(ThirdPartyWarehouse.name.in_(valid_warehouse_names))
        )
    ).scalars().all()
    warehouse_by_name = {row.name: row for row in existing_rows}
    new_warehouses = [name for name in valid_warehouse_names if name not in warehouse_by_name]
    unmaintained_warehouses = sorted(
        {
            name
            for name in valid_warehouse_names
            if name not in warehouse_by_name or warehouse_by_name[name].country is None
        }
    )
    issues = [
        {
            "row": row.source_row_no,
            "warehouse_name": row.warehouse_name or None,
            "commodity_sku": row.commodity_sku,
            "message": row.error_message,
        }
        for row in parsed_rows
        if row.error_message is not None
    ]

    batch = ThirdPartyInventoryImportBatch(
        filename=filename[:255] or "third-party-inventory.xlsx",
        status="pending",
        row_count=len(parsed_rows),
        valid_row_count=len(valid_rows),
        skipped_row_count=len(parsed_rows) - len(valid_rows),
        new_warehouse_count=len(new_warehouses),
        created_by=created_by,
        summary={
            "new_warehouses": new_warehouses,
            "unmaintained_warehouses": unmaintained_warehouses,
            "issues": issues[:PREVIEW_ISSUE_LIMIT],
            "issue_count": len(issues),
        },
    )
    db.add(batch)
    await db.flush()

    db.add_all(
        [
            ThirdPartyInventoryImportItem(
                batch_id=batch.id,
                warehouse_name_raw=row.warehouse_name,
                warehouse_id=warehouse_by_name.get(row.warehouse_name).id
                if row.warehouse_name in warehouse_by_name
                else None,
                commodity_sku=row.commodity_sku,
                available=row.available,
                reserved=row.reserved,
                source_row_no=row.source_row_no,
                error_message=row.error_message,
            )
            for row in parsed_rows
        ]
    )
    await db.commit()
    await db.refresh(batch)
    return batch


async def confirm_import_batch(
    db: AsyncSession,
    *,
    batch_id: int,
    confirmed_by: str | None,
) -> ThirdPartyInventoryImportBatch:
    batch = (
        await db.execute(
            select(ThirdPartyInventoryImportBatch).where(
                ThirdPartyInventoryImportBatch.id == batch_id
            )
        )
    ).scalar_one_or_none()
    if batch is None:
        raise NotFound("导入批次不存在")
    if batch.status != "pending":
        raise ConflictError("只有待确认批次可以确认导入")

    item_rows = (
        await db.execute(
            select(ThirdPartyInventoryImportItem)
            .where(ThirdPartyInventoryImportItem.batch_id == batch_id)
            .order_by(ThirdPartyInventoryImportItem.source_row_no)
        )
    ).scalars().all()
    valid_items = [
        item
        for item in item_rows
        if item.error_message is None
        and item.warehouse_name_raw
        and item.commodity_sku
        and item.available is not None
        and item.reserved is not None
    ]
    if not valid_items:
        raise ValidationFailed("导入批次没有有效库存行")

    names = sorted({item.warehouse_name_raw for item in valid_items})
    existing_rows = (
        await db.execute(select(ThirdPartyWarehouse).where(ThirdPartyWarehouse.name.in_(names)))
    ).scalars().all()
    warehouse_by_name = {row.name: row for row in existing_rows}
    for name in names:
        if name not in warehouse_by_name:
            warehouse = ThirdPartyWarehouse(name=name, country=None)
            db.add(warehouse)
            await db.flush()
            warehouse_by_name[name] = warehouse

    for item in valid_items:
        warehouse = warehouse_by_name[item.warehouse_name_raw]
        if item.warehouse_id != warehouse.id:
            item.warehouse_id = warehouse.id

    totals: defaultdict[tuple[int, str], dict[str, int]] = defaultdict(
        lambda: {"available": 0, "reserved": 0}
    )
    for item in valid_items:
        warehouse = warehouse_by_name[item.warehouse_name_raw]
        key = (warehouse.id, item.commodity_sku or "")
        totals[key]["available"] += int(item.available or 0)
        totals[key]["reserved"] += int(item.reserved or 0)

    await db.execute(delete(ThirdPartyInventoryCurrent))
    insert_values = [
        {
            "warehouse_id": warehouse_id,
            "commodity_sku": commodity_sku,
            "available": values["available"],
            "reserved": values["reserved"],
            "source_batch_id": batch.id,
            "last_operation": "import",
        }
        for (warehouse_id, commodity_sku), values in sorted(totals.items())
    ]
    if insert_values:
        await db.execute(pg_insert(ThirdPartyInventoryCurrent).values(insert_values))

    batch.status = "applied"
    batch.confirmed_by = confirmed_by
    batch.confirmed_at = datetime.now(BEIJING)
    await db.commit()
    await db.refresh(batch)
    return batch


async def expire_pending_batch(db: AsyncSession, batch_id: int) -> ThirdPartyInventoryImportBatch:
    batch = (
        await db.execute(
            select(ThirdPartyInventoryImportBatch).where(
                ThirdPartyInventoryImportBatch.id == batch_id
            )
        )
    ).scalar_one_or_none()
    if batch is None:
        raise NotFound("导入批次不存在")
    if batch.status != "pending":
        raise ConflictError("只有待确认批次可以取消")
    batch.status = "expired"
    await db.commit()
    await db.refresh(batch)
    return batch


async def upsert_current_item(
    db: AsyncSession,
    *,
    warehouse_id: int,
    commodity_sku: str,
    available: int,
    reserved: int,
) -> ThirdPartyInventoryCurrent:
    warehouse = (
        await db.execute(select(ThirdPartyWarehouse).where(ThirdPartyWarehouse.id == warehouse_id))
    ).scalar_one_or_none()
    if warehouse is None:
        raise NotFound("三方仓不存在")
    stmt = (
        pg_insert(ThirdPartyInventoryCurrent)
        .values(
            warehouse_id=warehouse_id,
            commodity_sku=commodity_sku,
            available=available,
            reserved=reserved,
            source_batch_id=None,
            last_operation="manual",
        )
        .on_conflict_do_update(
            constraint="uq_third_party_inventory_current_wh_sku",
            set_={
                "available": available,
                "reserved": reserved,
                "source_batch_id": None,
                "last_operation": "manual",
                "updated_at": func.now(),
            },
        )
        .returning(ThirdPartyInventoryCurrent)
    )
    item = (await db.execute(stmt)).scalar_one()
    await db.commit()
    return item


async def patch_current_item(
    db: AsyncSession,
    *,
    item_id: int,
    values: dict[str, Any],
) -> ThirdPartyInventoryCurrent:
    item = (
        await db.execute(
            select(ThirdPartyInventoryCurrent).where(ThirdPartyInventoryCurrent.id == item_id)
        )
    ).scalar_one_or_none()
    if item is None:
        raise NotFound("三方库存明细不存在")
    warehouse_id = int(values.get("warehouse_id") or item.warehouse_id)
    warehouse = (
        await db.execute(select(ThirdPartyWarehouse).where(ThirdPartyWarehouse.id == warehouse_id))
    ).scalar_one_or_none()
    if warehouse is None:
        raise NotFound("三方仓不存在")
    item.warehouse_id = warehouse_id
    item.commodity_sku = values.get("commodity_sku") or item.commodity_sku
    if values.get("available") is not None:
        item.available = int(values["available"])
    if values.get("reserved") is not None:
        item.reserved = int(values["reserved"])
    item.source_batch_id = None
    item.last_operation = "manual"
    try:
        await db.commit()
    except Exception as exc:
        await db.rollback()
        raise ConflictError("同一三方仓下 SKU 已存在") from exc
    await db.refresh(item)
    return item


async def delete_current_item(db: AsyncSession, item_id: int) -> None:
    result = await db.execute(
        delete(ThirdPartyInventoryCurrent).where(ThirdPartyInventoryCurrent.id == item_id)
    )
    if result.rowcount == 0:
        raise NotFound("三方库存明细不存在")
    await db.commit()


async def sync_import_items_warehouse_id(db: AsyncSession, warehouse: ThirdPartyWarehouse) -> None:
    await db.execute(
        update(ThirdPartyInventoryImportItem)
        .where(
            ThirdPartyInventoryImportItem.warehouse_name_raw == warehouse.name,
            ThirdPartyInventoryImportItem.warehouse_id.is_(None),
            ThirdPartyInventoryImportItem.error_message.is_(None),
        )
        .values(warehouse_id=warehouse.id)
    )
